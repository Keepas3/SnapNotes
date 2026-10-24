"""Turn uploaded files into Claude content blocks.

PDFs and images are sent natively (Claude reads text and page imagery, so no separate OCR step).
DOCX/PPTX/Markdown/plain text are parsed locally into delimited text.
"""
from __future__ import annotations

import base64
from dataclasses import dataclass
from pathlib import Path
from typing import Sequence

from app.core.config import settings

IMAGE_TYPES = {
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".png": "image/png",
    ".gif": "image/gif",
    ".webp": "image/webp",
}
TEXT_SUFFIXES = {".txt", ".md", ".markdown"}
MAX_IMAGE_BYTES = 5 * 1024 * 1024  # per-image API limit


class UnsupportedFileError(ValueError):
    pass


@dataclass
class SourceDoc:
    name: str
    blocks: list[dict]  # Claude content blocks, without cache_control

    def content(self, tail_text: str) -> list[dict]:
        """Notes first (cache breakpoint on the last note block), stage instructions last."""
        blocks = [dict(b) for b in self.blocks]
        blocks[-1] = {**blocks[-1], "cache_control": {"type": "ephemeral"}}
        return blocks + [{"type": "text", "text": tail_text}]


def _wrap_notes(filename: str, body: str) -> dict:
    # Stop note text from closing the delimiter early (cheap prompt-injection hardening).
    safe = body.replace("</notes>", "<\\/notes>")
    return {"type": "text", "text": f'<notes file="{filename}">\n{safe}\n</notes>'}


def _docx_text(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    lines: list[str] = []
    for p in doc.paragraphs:
        if not p.text.strip():
            continue
        style = (p.style.name or "") if p.style is not None else ""
        if style.startswith("Heading"):
            level = "".join(ch for ch in style if ch.isdigit()) or "1"
            lines.append(f"{'#' * int(level)} {p.text.strip()}")
        else:
            lines.append(p.text.strip())
    for t_index, table in enumerate(doc.tables, 1):
        lines.append(f"[Table {t_index}]")
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)


def _pptx_text(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(path))
    out: list[str] = []
    for i, slide in enumerate(prs.slides, 1):
        out.append(f"--- Slide {i} ---")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                out.append(shape.text_frame.text.strip())
            elif getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    out.append(" | ".join(cell.text.strip() for cell in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            out.append(f"[Speaker notes] {slide.notes_slide.notes_text_frame.text.strip()}")
    return "\n".join(out)


def _file_blocks(path: Path) -> list[dict]:
    suffix = path.suffix.lower()
    size = path.stat().st_size
    if size > settings.max_upload_bytes:
        raise UnsupportedFileError(f"{path.name} is {size / 1e6:.1f} MB; the limit is {settings.max_upload_bytes / 1e6:.0f} MB")

    if suffix == ".pdf":
        data = base64.standard_b64encode(path.read_bytes()).decode("ascii")
        return [
            {"type": "text", "text": f'Attached PDF: "{path.name}"'},
            {"type": "document", "source": {"type": "base64", "media_type": "application/pdf", "data": data}},
        ]
    if suffix in IMAGE_TYPES:
        if size > MAX_IMAGE_BYTES:
            raise UnsupportedFileError(f"{path.name} is over 5 MB; resize or compress the photo first")
        data = base64.standard_b64encode(path.read_bytes()).decode("ascii")
        return [
            {"type": "text", "text": f'Photo of notes: "{path.name}"'},
            {"type": "image", "source": {"type": "base64", "media_type": IMAGE_TYPES[suffix], "data": data}},
        ]
    if suffix == ".docx":
        return [_wrap_notes(path.name, _docx_text(path))]
    if suffix == ".pptx":
        return [_wrap_notes(path.name, _pptx_text(path))]
    if suffix in TEXT_SUFFIXES:
        return [_wrap_notes(path.name, path.read_text(encoding="utf-8", errors="replace"))]
    raise UnsupportedFileError(f"Unsupported file type: {path.name}")


def load_source(paths: Sequence[str | Path]) -> SourceDoc:
    """One document can be several files, e.g. several photos of consecutive notebook pages."""
    if not paths:
        raise UnsupportedFileError("No files given")
    resolved = [Path(p) for p in paths]
    blocks: list[dict] = []
    for p in resolved:
        blocks.extend(_file_blocks(p))
    name = resolved[0].name if len(resolved) == 1 else f"{resolved[0].name} (+{len(resolved) - 1} more)"
    return SourceDoc(name=name, blocks=blocks)


def validate_file(path: Path) -> None:
    """Cheap upload-time check, without base64-encoding large files: type, size, magic bytes, parseability."""
    suffix = path.suffix.lower()
    size = path.stat().st_size
    if size == 0:
        raise UnsupportedFileError(f"{path.name} is empty")
    if size > settings.max_upload_bytes:
        raise UnsupportedFileError(f"{path.name} is {size / 1e6:.1f} MB; the limit is {settings.max_upload_bytes / 1e6:.0f} MB")
    if suffix == ".pdf":
        with path.open("rb") as fh:
            if not fh.read(5).startswith(b"%PDF-"):
                raise UnsupportedFileError(f"{path.name} is not a valid PDF")
    elif suffix in IMAGE_TYPES:
        if size > MAX_IMAGE_BYTES:
            raise UnsupportedFileError(f"{path.name} is over 5 MB; resize or compress the photo first")
    elif suffix == ".docx":
        _docx_text(path)
    elif suffix == ".pptx":
        _pptx_text(path)
    elif suffix not in TEXT_SUFFIXES:
        raise UnsupportedFileError(f"Unsupported file type: {path.name}")


ALLOWED_SUFFIXES = frozenset({".pdf", ".docx", ".pptx", *IMAGE_TYPES, *TEXT_SUFFIXES})
